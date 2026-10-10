"""Fixed, bounded Office reader; this module also runs inside the workspace image."""

from __future__ import annotations

import hashlib
import io
import json
import os
import posixpath
import re
import stat
import sys
from typing import Any

from defusedxml import ElementTree
from defusedxml.common import DefusedXmlException

from runtime_service.workspace.archives import read_zip

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
OFFICE_MIMES = {DOCX_MIME: "docx", PPTX_MIME: "pptx"}
MAX_CHARS = 12_000
MAX_PARTS = 20
MAX_BYTES = 20 * 1024 * 1024
_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_OFFICE_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def validate_office(data: bytes, mime: str) -> list[str]:
    """Validate the package without calling a document parser or following links."""
    code = "invalid_document" if mime == DOCX_MIME else "invalid_presentation"
    if mime not in OFFICE_MIMES or not data or len(data) > MAX_BYTES:
        raise ValueError(code)
    try:
        members = dict(read_zip(data))
        required = "word/document.xml" if mime == DOCX_MIME else "ppt/presentation.xml"
        if not {"[Content_Types].xml", "_rels/.rels", required} <= members.keys():
            raise ValueError(code)
        _validate_main_part(members, required, mime)
        external = False
        for name, raw in members.items():
            lowered = name.lower()
            if any(
                part in lowered for part in ("vbaproject", "/embeddings/", "activex")
            ):
                raise ValueError(code)
            if not lowered.endswith((".xml", ".rels")):
                continue
            root = ElementTree.fromstring(raw, forbid_dtd=True)
            if name == "[Content_Types].xml":
                for item in root:
                    kind = item.get("ContentType", "").lower()
                    if any(
                        word in kind
                        for word in ("macroenabled", "vba", "oleobject", "activex")
                    ):
                        raise ValueError(code)
            if lowered.endswith(".rels"):
                external |= _validate_relationships(name, root, members)
        if not _has_main_relationship(members, required):
            raise ValueError(code)
        return ["external_relationship_ignored"] if external else []
    except (
        ValueError,
        KeyError,
        ElementTree.ParseError,
        DefusedXmlException,
        OSError,
    ) as exc:
        raise ValueError(code) from exc


def _validate_main_part(members: dict[str, bytes], main: str, mime: str) -> None:
    types = ElementTree.fromstring(members["[Content_Types].xml"], forbid_dtd=True)
    prefix = (
        "wordprocessingml.document"
        if mime == DOCX_MIME
        else "presentationml.presentation"
    )
    expected_type = f"application/vnd.openxmlformats-officedocument.{prefix}.main+xml"
    if not any(
        item.get("PartName") == "/" + main and item.get("ContentType") == expected_type
        for item in types
    ):
        raise ValueError("invalid_main_part")
    root = ElementTree.fromstring(members[main], forbid_dtd=True)
    expected_tag = (
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}document"
        if mime == DOCX_MIME
        else "{http://schemas.openxmlformats.org/presentationml/2006/main}presentation"
    )
    if root.tag != expected_tag:
        raise ValueError("invalid_main_part")
    if mime == PPTX_MIME:
        relations = ElementTree.fromstring(
            members["ppt/_rels/presentation.xml.rels"], forbid_dtd=True
        )
        slide_ids = {
            item.get("Id")
            for item in relations
            if item.get("Type") == f"{_OFFICE_REL}/slide"
        }
        slide_tag = "{http://schemas.openxmlformats.org/presentationml/2006/main}sldId"
        if any(
            item.get(f"{{{_OFFICE_REL}}}id") not in slide_ids
            for item in root.iter(slide_tag)
        ):
            raise ValueError("invalid_slide_relationship")


def _validate_relationships(name: str, root, members: dict[str, bytes]) -> bool:
    if root.tag != f"{{{_REL}}}Relationships":
        raise ValueError("invalid_relationships")
    base = "" if name == "_rels/.rels" else posixpath.dirname(posixpath.dirname(name))
    seen: set[str] = set()
    external = False
    for relation in root:
        target = relation.get("Target", "")
        identity = relation.get("Id", "")
        kind = relation.get("Type", "")
        if (
            not identity
            or identity in seen
            or not target
            or relation.tag != f"{{{_REL}}}Relationship"
        ):
            raise ValueError("invalid_relationships")
        seen.add(identity)
        if relation.get("TargetMode") == "External":
            if kind != f"{_OFFICE_REL}/hyperlink":
                raise ValueError("external_relationship_denied")
            external = True
            continue
        if relation.get("TargetMode") not in (None, "Internal"):
            raise ValueError("invalid_relationships")
        if (
            "\\" in target
            or ":" in target
            or "%" in target
            or any(ord(c) < 32 for c in target)
        ):
            raise ValueError("invalid_relationships")
        resolved = (
            posixpath.normpath(posixpath.join(base, target.lstrip("/")))
            if not target.startswith("/")
            else target.lstrip("/")
        )
        if resolved.startswith("../") or resolved not in members:
            raise ValueError("invalid_relationships")
    return external


def _has_main_relationship(members: dict[str, bytes], main: str) -> bool:
    root = ElementTree.fromstring(members["_rels/.rels"], forbid_dtd=True)
    return any(
        item.get("Type") == f"{_OFFICE_REL}/officeDocument"
        and item.get("Target", "").lstrip("/") == main
        and item.get("TargetMode") != "External"
        for item in root
    )


def validate_read_options(
    mime: str,
    query: str | None,
    pages: tuple[int | None, int | None],
    options: dict | None,
) -> dict[str, int]:
    if query is not None and (
        not isinstance(query, str) or not query.strip() or len(query) > 500
    ):
        raise ValueError("invalid_query")
    if options is not None and not isinstance(options, dict):
        raise ValueError("invalid_read_options")
    result = options or {}
    allowed = (
        {"section_start", "section_end", "char_offset"}
        if mime == DOCX_MIME
        else {"char_offset"}
    )
    if set(result) - allowed or any(
        type(value) is not int for value in result.values()
    ):
        raise ValueError("invalid_read_options")
    if result.get("char_offset", 0) < 0:
        raise ValueError("invalid_read_options")
    if any(
        value is not None and (type(value) is not int or value < 1) for value in pages
    ):
        raise ValueError("invalid_page_range")
    if mime == DOCX_MIME and any(value is not None for value in pages):
        raise ValueError("invalid_page_range")
    if any(result.get(key, 1) < 1 for key in ("section_start", "section_end")):
        raise ValueError("invalid_read_options")
    return result


def _docx_parts(data: bytes) -> list[str]:
    from docx import Document
    from docx.table import Table

    document = Document(io.BytesIO(data))
    return [
        "\n".join("\t".join(cell.text for cell in row.cells) for row in block.rows)
        if isinstance(block, Table)
        else block.text
        for block in document.iter_inner_content()
    ]


def _pptx_parts(data: bytes) -> list[str]:
    from pptx import Presentation

    def texts(shapes):
        for shape in shapes:
            if hasattr(shape, "shapes"):
                yield from texts(shape.shapes)
            elif shape.has_text_frame:
                yield shape.text_frame.text
            elif shape.has_table:
                yield "\n".join(
                    "\t".join(cell.text for cell in row.cells)
                    for row in shape.table.rows
                )

    return [
        "\n".join(texts(slide.shapes))
        for slide in Presentation(io.BytesIO(data)).slides
    ]


def read_office(data: bytes, mime: str, request: dict[str, Any]) -> dict[str, Any]:
    options = validate_read_options(
        mime,
        request.get("query"),
        (request.get("page_start"), request.get("page_end")),
        request.get("read_options"),
    )
    warnings = validate_office(data, mime)
    try:
        parts = _docx_parts(data) if mime == DOCX_MIME else _pptx_parts(data)
    except (
        ValueError,
        KeyError,
        TypeError,
        AttributeError,
        OSError,
        RecursionError,
        ElementTree.ParseError,
    ) as exc:
        raise ValueError(
            "invalid_document" if mime == DOCX_MIME else "invalid_presentation"
        ) from exc
    if mime == DOCX_MIME:
        warnings += ["docx_body_only", "docx_table_structure_flattened"]
    else:
        warnings += ["presentation_text_only"]
    return _select_parts(parts, mime, request, options, warnings)


def _select_parts(
    parts: list[str], mime: str, request: dict, options: dict, warnings: list[str]
) -> dict:
    docx = mime == DOCX_MIME
    start = options.get("section_start", 1) if docx else request.get("page_start") or 1
    end = (
        options.get("section_end", min(len(parts), start + MAX_PARTS - 1))
        if docx
        else request.get("page_end") or min(len(parts), start + MAX_PARTS - 1)
    )
    if not parts and start == 1 and end == 0 and not options.get("char_offset"):
        return {
            "version": 1,
            "format": OFFICE_MIMES[mime],
            "sections" if docx else "pages": 0,
            "read_range": [],
            "matched_sections" if docx else "matched_pages": [],
            "text": "",
            "chunks": [],
            "truncated": False,
            "warnings": [*warnings, "document_no_text"],
            "next_read": None,
        }
    if start < 1 or end < start or end > len(parts) or end - start + 1 > MAX_PARTS:
        raise ValueError("invalid_read_options" if docx else "invalid_page_range")
    query = request.get("query")
    matches = [
        (index + 1, text)
        for index, text in enumerate(parts[start - 1 : end], start - 1)
        if query is None or query.casefold() in text.casefold()
    ]
    if query and not matches:
        warnings.append("no_query_match_in_selected_range")
    if docx and not any(text.strip() for text in parts[start - 1 : end]):
        warnings.append("document_no_text")
    if not docx:
        warnings.extend(
            f"slide_{index}_no_text_layer_ocr_required"
            for index in range(start, end + 1)
            if not parts[index - 1].strip()
        )
    offset = options.get("char_offset", 0)
    if offset and (not matches or offset >= len(matches[0][1])):
        raise ValueError("invalid_read_options")
    chunks, next_read = _bounded_chunks(matches, docx, request, offset)
    if next_read is None and end < len(parts):
        next_read = _next_request(request, docx, end + 1)
    return {
        "version": 1,
        "format": OFFICE_MIMES[mime],
        "sections" if docx else "pages": len(parts),
        "read_range": [start, end],
        "matched_sections" if docx else "matched_pages": [
            index for index, _ in matches
        ],
        "text": "\n\n".join(chunk["text"] for chunk in chunks),
        "chunks": chunks,
        "truncated": next_read is not None,
        "warnings": warnings,
        "next_read": next_read,
    }


def _next_request(
    request: dict, docx: bool, number: int, offset: int = 0, end: int | None = None
) -> dict:
    result = {"file_path": request["file_path"]}
    if request.get("query") is not None:
        result["query"] = request["query"]
    options = {"char_offset": offset} if offset else {}
    if docx:
        options["section_start"] = number
        if end is not None:
            options["section_end"] = end
    else:
        result["page_start"] = number
        if end is not None:
            result["page_end"] = end
    if options:
        result["read_options"] = options
    return result


def _bounded_chunks(
    matches: list[tuple[int, str]], docx: bool, request: dict, offset: int
) -> tuple[list[dict], dict | None]:
    chunks: list[dict] = []
    remaining = MAX_CHARS
    for index, (number, text) in enumerate(matches):
        begin = offset if index == 0 else 0
        separator = 2 if chunks else 0
        available = max(0, remaining - separator)
        chunk = text[begin : begin + available]
        if chunk:
            chunks.append(
                {
                    "section" if docx else "page": number,
                    "char_offset": begin,
                    "text": chunk,
                }
            )
            remaining -= len(chunk) + separator
        if begin + len(chunk) < len(text):
            end = (
                (request.get("read_options") or {}).get("section_end")
                if docx
                else request.get("page_end")
            )
            return chunks, _next_request(request, docx, number, begin + len(chunk), end)
    return chunks, None


def _read_source(request: dict) -> tuple[bytes, str]:
    path = request.get("file_path", "")
    match = (
        re.fullmatch(r"/workspace/(uploads|outputs)/([0-9a-f]{64})\.(docx|pptx)", path)
        if isinstance(path, str)
        else None
    )
    if match is None:
        raise ValueError("invalid_file_ref")
    folder, digest, extension = match.groups()
    # Open relative to descriptors: do not follow a swapped directory or source symlink.
    root = os.open("/workspace", os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW)
    directory = None
    try:
        directory = os.open(
            folder, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW, dir_fd=root
        )
        descriptor = os.open(
            f"{digest}.{extension}",
            os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK,
            dir_fd=directory,
        )
        with os.fdopen(descriptor, "rb") as source:
            if not stat.S_ISREG(os.fstat(source.fileno()).st_mode):
                raise ValueError("invalid_file_ref")
            raw = source.read(MAX_BYTES + 1)
    finally:
        if directory is not None:
            os.close(directory)
        os.close(root)
    if len(raw) > MAX_BYTES or hashlib.sha256(raw).hexdigest() != digest:
        raise ValueError("file_hash_mismatch")
    return raw, next(mime for mime, ext in OFFICE_MIMES.items() if ext == extension)


def main() -> int:
    try:
        request = json.loads(sys.argv[1])
        if not isinstance(request, dict) or set(request) - {
            "file_path",
            "query",
            "page_start",
            "page_end",
            "read_options",
        }:
            raise ValueError("invalid_read_options")
        data, mime = _read_source(request)
        result = read_office(data, mime, request)
    except (ValueError, OSError) as exc:
        allowed = {
            "invalid_file_ref",
            "file_hash_mismatch",
            "invalid_query",
            "invalid_page_range",
            "invalid_read_options",
            "invalid_document",
            "invalid_presentation",
        }
        code = str(exc) if str(exc) in allowed else "office_read_failed"
        result = {"error": code}
    print(json.dumps(result, ensure_ascii=False, separators=(",", ":")))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
