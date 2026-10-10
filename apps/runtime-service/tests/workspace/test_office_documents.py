import hashlib
import io
import zipfile

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from runtime_service.workspace.document_reader import DOCX_MIME, PPTX_MIME, read_office
from runtime_service.workspace.documents import (
    DocumentError,
    DocumentWorkspace,
    validate_document,
)


def docx_bytes(*paragraphs, table=False):
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        row = document.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text, row.cells[1].text = "Revenue", "125"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def pptx_bytes(*slides):
    presentation = Presentation()
    for text in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        if text:
            slide.shapes.add_textbox(
                Inches(1), Inches(1), Inches(5), Inches(2)
            ).text = text
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def replace_member(data, name, replacement):
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as source,
        zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target,
    ):
        for item in source.infolist():
            if item.filename != name:
                target.writestr(item.filename, source.read(item))
        if replacement is not None:
            target.writestr(name, replacement)
    return output.getvalue()


@pytest.mark.parametrize(
    "mime,raw",
    [
        (DOCX_MIME, docx_bytes("合同 Contract", table=True)),
        (PPTX_MIME, pptx_bytes("Forecast 125", "")),
    ],
    ids=["docx", "pptx"],
)
def test_office_upload_and_read(tmp_path, mime, raw):
    digest = hashlib.sha256(raw).hexdigest()
    ref = DocumentWorkspace(tmp_path).put(
        raw, digest, mime, "source.docx" if mime == DOCX_MIME else "source.pptx"
    )
    assert set(ref) == {
        "version",
        "path",
        "file_name",
        "mime_type",
        "size_bytes",
        "sha256",
    }
    result = read_office(raw, mime, {"file_path": ref["path"]})
    assert "125" in result["text"]
    assert result["read_range"][0] == 1 and result["next_read"] is None
    assert (
        "Contract" in result["text"]
        if mime == DOCX_MIME
        else "slide_2_no_text_layer_ocr_required" in result["warnings"]
    )


@pytest.mark.parametrize(
    "mime,builder", [(DOCX_MIME, docx_bytes), (PPTX_MIME, pptx_bytes)]
)
def test_long_part_can_be_completely_resumed(mime, builder):
    text = "".join(f"Value {index:05} 中文; " for index in range(2000))
    raw = builder(text, "Second part")
    request = {
        "file_path": "/workspace/uploads/"
        + "a" * 64
        + (".docx" if mime == DOCX_MIME else ".pptx")
    }
    request.update(
        {"read_options": {"section_end": 1}} if mime == DOCX_MIME else {"page_end": 1}
    )
    chunks = []
    while request is not None:
        result = read_office(raw, mime, request)
        assert len(result["text"]) <= 12000
        chunks.append(result["text"])
        request = result["next_read"]
        if request is not None and (
            request.get("page_start", 1) == 2
            or request.get("read_options", {}).get("section_start") == 2
        ):
            break
    assert "".join(chunks) == text


def test_slides_range_query_and_empty_document():
    raw = pptx_bytes(*(f"Slide {index}" for index in range(1, 25)))
    request = {
        "file_path": "/workspace/uploads/" + "a" * 64 + ".pptx",
        "query": "Slide 24",
    }
    result = read_office(raw, PPTX_MIME, request)
    assert result["text"] == "" and result["next_read"]["page_start"] == 21
    assert read_office(raw, PPTX_MIME, result["next_read"])["matched_pages"] == [24]
    result = read_office(docx_bytes(), DOCX_MIME, {"file_path": "input"})
    assert result["sections"] == 0 and "document_no_text" in result["warnings"]
    result = read_office(docx_bytes(""), DOCX_MIME, {"file_path": "input"})
    assert result["text"] == "" and "document_no_text" in result["warnings"]


@pytest.mark.parametrize(
    "patch",
    [
        ("word/vbaProject.bin", b"macro"),
        ("word/embeddings/oleObject1.bin", b"embedded"),
        ("../escape", b"escape"),
        ("word/document.xml", b'<!DOCTYPE x [<!ENTITY bad "boom">]><x>&bad;</x>'),
        ("word/document.xml", b"broken"),
        ("word/document.xml", None),
        (
            "word/_rels/document.xml.rels",
            b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate" Target="https://example.com/template" TargetMode="External"/></Relationships>',
        ),
        (
            "word/_rels/document.xml.rels",
            b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="x" Target="missing.xml"/></Relationships>',
        ),
        ("large.bin", b"0" * (2 * 1024 * 1024 + 1)),
    ],
    ids=[
        "macro",
        "ole",
        "traversal",
        "entity",
        "broken",
        "missing",
        "template",
        "missing-target",
        "expansion",
    ],
)
def test_dangerous_office_is_rejected_without_storing(tmp_path, patch):
    raw = replace_member(docx_bytes("Test"), *patch)
    with pytest.raises(DocumentError, match="invalid_document"):
        DocumentWorkspace(tmp_path).put(raw, hashlib.sha256(raw).hexdigest(), DOCX_MIME)
    assert not (tmp_path / "uploads").exists()


@pytest.mark.parametrize(
    "options",
    [
        {"char_offset": -1},
        {"char_offset": True},
        {"section_start": 0},
        {"unknown": 1},
        {"section_start": 2, "section_end": 1},
    ],
)
def test_read_options_are_strict(options):
    with pytest.raises(ValueError):
        read_office(
            docx_bytes("Test"),
            DOCX_MIME,
            {"file_path": "source", "read_options": options},
        )


def test_hyperlink_is_not_followed():
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    paragraph = document.add_paragraph("Text: ")
    identity = paragraph.part.relate_to(
        "https://127.0.0.1/private", RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), identity)
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "Public link label"
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    output = io.BytesIO()
    document.save(output)
    result = read_office(output.getvalue(), DOCX_MIME, {"file_path": "source"})
    assert "external_relationship_ignored" in result["warnings"]
    assert "https://" not in result["text"]
    assert "Public link label" in result["text"]


def test_pptx_invalid_relationships_and_legacy_magic():
    raw = replace_member(
        pptx_bytes("Text"), "ppt/_rels/presentation.xml.rels", b"broken"
    )
    with pytest.raises(DocumentError, match="invalid_presentation"):
        validate_document(raw, PPTX_MIME)
    with pytest.raises(DocumentError):
        validate_document(bytes.fromhex("d0cf11e0a1b11ae1"), DOCX_MIME)


def test_office_package_count_and_content_type_are_bounded():
    raw = docx_bytes("count")
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        with zipfile.ZipFile(io.BytesIO(raw)) as source:
            for name in source.namelist():
                archive.writestr(name, source.read(name))
        for index in range(256):
            archive.writestr(f"extra-{index}.txt", b"data")
    with pytest.raises(DocumentError, match="invalid_document"):
        validate_document(output.getvalue(), DOCX_MIME)
    with zipfile.ZipFile(io.BytesIO(raw)) as source:
        types = source.read("[Content_Types].xml")
    wrong = replace_member(
        raw,
        "[Content_Types].xml",
        types.replace(
            b"wordprocessingml.document.main+xml",
            b"wordprocessingml.document.macroEnabled.main+xml",
        ),
    )
    with pytest.raises(DocumentError, match="invalid_document"):
        validate_document(wrong, DOCX_MIME)
    with pytest.raises(DocumentError, match="invalid_presentation"):
        validate_document(raw, PPTX_MIME)
